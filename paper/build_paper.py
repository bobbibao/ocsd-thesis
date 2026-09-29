# Build OCSD_paper.docx from paper.md: number citations, append references, convert with pandoc, then style.
import os, re, sys
HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, os.path.dirname(HERE))
import pypandoc
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from refs import REFS

REF = {r['key']: r['text'] for r in REFS}
REF['zhang2025sketchscene'] = ('Tianyu Zhang, Xiaoxuan Xie, Xusheng Du, Haoran Xie (2025), "Sketch-Guided Scene Image Generation '
                               'with Diffusion Model", *Computers & Graphics* 129, Article 104226 (preprint: arXiv:2407.06469, 2024).')
REF['mitsouras2024usketch'] = REF['mitsouras2024usketch'].replace('mã nguồn', 'code')

md = open(os.path.join(HERE, 'paper.md'), encoding='utf-8').read()
order = []
def cite(m):
    keys = [k.strip().lstrip('@') for k in m.group(1).split(';')]
    for k in keys:
        if k not in REF:
            raise SystemExit('unknown ref ' + k)
        if k not in order:
            order.append(k)
    nums = sorted(order.index(k) + 1 for k in keys)
    return '[' + ', '.join(map(str, nums)) + ']'
md = re.sub(r'\[(@[^\]]+)\]', cite, md)
md += '\n' + '\n\n'.join('::: {custom-style="Reference"}\n[%d] %s\n:::' % (i + 1, REF[k]) for i, k in enumerate(order)) + '\n'
src = os.path.join(HERE, '_paper_numbered.md')
open(src, 'w', encoding='utf-8').write(md)

out = os.path.join(HERE, 'OCSD_paper.docx')
pypandoc.convert_file(src, 'docx', outputfile=out, extra_args=['--resource-path', HERE],
                      format='markdown+pipe_tables+tex_math_dollars+fenced_divs+implicit_figures+subscript')

# ---------------------------------------------------------------- styling
doc = Document(out)
FONT = 'Times New Roman'
sec = doc.sections[0]
sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
sec.left_margin = sec.right_margin = Cm(2.2)
sec.top_margin = sec.bottom_margin = Cm(2.2)

def font(style, size=None, bold=None, italic=None, color=None):
    f = style.font
    f.name = FONT
    rpr = style.element.get_or_add_rPr()
    rf = rpr.find(qn('w:rFonts'))
    if rf is None:
        rf = OxmlElement('w:rFonts'); rpr.append(rf)
    for a in ('w:ascii', 'w:hAnsi', 'w:cs', 'w:eastAsia'):
        rf.set(qn(a), FONT)
    for a in ('w:asciiTheme', 'w:hAnsiTheme', 'w:cstheme', 'w:eastAsiaTheme'):
        if rf.get(qn(a)) is not None:
            del rf.attrib[qn(a)]
    if size: f.size = Pt(size)
    if bold is not None: f.bold = bold
    if italic is not None: f.italic = italic
    f.color.rgb = RGBColor.from_string(color or '000000')

S = doc.styles
def para(name, size, align=None, before=0, after=4, line=1.1, bold=None, italic=None, indent=None, left=None, right=None):
    st = S[name]
    font(st, size, bold, italic)
    pf = st.paragraph_format
    if align is not None: pf.alignment = align
    pf.space_before, pf.space_after, pf.line_spacing = Pt(before), Pt(after), line
    if indent is not None: pf.first_line_indent = indent
    if left is not None: pf.left_indent = left
    if right is not None: pf.right_indent = right

J, C = WD_ALIGN_PARAGRAPH.JUSTIFY, WD_ALIGN_PARAGRAPH.CENTER
for n in ('Normal', 'Body Text', 'First Paragraph', 'Compact'):
    if n in [s.name for s in S]:
        para(n, 10.5, J, 0, 5, 1.12)
para('Title', 17, C, 0, 8, 1.0, bold=True)
for n, sz, b in (('Heading 1', 12.5, 12), ('Heading 2', 11, 8)):
    para(n, sz, WD_ALIGN_PARAGRAPH.LEFT, b, 4, 1.0, bold=True)
para('Author', 11.5, C, 0, 2, 1.0)
para('Affiliation', 10, C, 0, 10, 1.0, italic=True)
para('Note', 9, J, 0, 10, 1.05, left=Cm(0.4), right=Cm(0.4))
para('AbstractTitle', 11, C, 4, 2, 1.0, bold=True)
para('Abstract', 9.8, J, 0, 8, 1.08, left=Cm(0.8), right=Cm(0.8))
para('TableCaption', 9.5, C, 8, 3, 1.0)
para('Image Caption', 9.5, J, 2, 10, 1.05, italic=False)
para('Captioned Figure', 10, C, 8, 0, 1.0)
para('Reference', 9, J, 0, 2, 1.0, indent=Cm(-0.7), left=Cm(0.7))
if 'Compact' in [s.name for s in S]:
    para('Compact', 10.5, J, 0, 2, 1.1)

# shade the status note
for p in doc.paragraphs:
    if p.style.name == 'Note':
        ppr = p._p.get_or_add_pPr()
        shd = OxmlElement('w:shd'); shd.set(qn('w:val'), 'clear'); shd.set(qn('w:fill'), 'FFF4D6'); ppr.append(shd)
    if p.style.name in ('Heading 1', 'Heading 2'):
        for r in p.runs:
            r.font.color.rgb = RGBColor(0, 0, 0)

# tables: three-line style, small font, centred
def border(el, tag, sz):
    b = OxmlElement('w:' + tag)
    b.set(qn('w:val'), 'single'); b.set(qn('w:sz'), str(sz)); b.set(qn('w:color'), '000000'); b.set(qn('w:space'), '0')
    el.append(b)
for t in doc.tables:
    t.alignment = 1
    t.autofit = False
    lay = OxmlElement('w:tblLayout'); lay.set(qn('w:type'), 'fixed'); t._tbl.tblPr.append(lay)
    tw = t._tbl.tblPr.find(qn('w:tblW'))
    if tw is not None:
        tw.set(qn('w:type'), 'dxa'); tw.set(qn('w:w'), str(int(Cm(16.6) / 635)))
    tblPr = t._tbl.tblPr
    for old in tblPr.findall(qn('w:tblBorders')):
        tblPr.remove(old)
    bs = OxmlElement('w:tblBorders'); border(bs, 'top', 12); border(bs, 'bottom', 12); tblPr.append(bs)
    ncol = len(t.columns)
    first = Cm(4.2) if ncol > 4 else Cm(3.0)
    rest = (Cm(16.6) - first) / (ncol - 1)
    widths = [first] + [int(rest)] * (ncol - 1)
    grid = t._tbl.find(qn('w:tblGrid'))
    for j, gc in enumerate(grid.findall(qn('w:gridCol'))):
        gc.set(qn('w:w'), str(int(widths[j] / 635)))
    for i, row in enumerate(t.rows):
        for j, cell in enumerate(row.cells):
            cell.width = widths[j]
            for p in cell.paragraphs:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT if j == 0 else WD_ALIGN_PARAGRAPH.CENTER
            if i == 0:
                tcPr = cell._tc.get_or_add_tcPr()
                tb = OxmlElement('w:tcBorders'); border(tb, 'bottom', 6); tcPr.append(tb)
            for p in cell.paragraphs:
                p.paragraph_format.space_before = Pt(1); p.paragraph_format.space_after = Pt(1)
                for r in p.runs:
                    r.font.size = Pt(8.3); r.font.name = FONT
                    if i == 0: r.font.bold = True
doc.core_properties.title = 'Object-Consistent Sketch-and-Text Guided Scene Image Generation with Diffusion Models'
doc.core_properties.author = 'Le Hoang Bao, Thai Ba Cuong, Nguyen Thanh Chuyen'
doc.save(out)
print('saved', out, len(order), 'references')
